# score_graphrag_ragas.py
#
# Reads graphrag_ragas_inputs.json (output of run_graphrag_eval.py)
# and runs RAGAS Faithfulness + Context Recall scoring on it.
#
# Judge model: openai/gpt-4o-mini via OpenRouter -- same judge, same
# settings as score_naive_rag_ragas.py, deliberately. RAGAS scores are
# only comparable across naive RAG and GraphRAG if the same judge scored
# both under identical conditions. Requires OPENROUTER_API_KEY.
#
# No pipeline runs here -- just scoring pre-generated outputs.
#
# Usage:
#   python score_graphrag_ragas.py
#
# Output:
#   ragas_outputs_graphrag/aeroops_graphrag_ragas_results.xlsx

from __future__ import annotations

import json
import math
import os
import re
import statistics
import sys
import time
from pathlib import Path
from typing import Any, Dict, List

# Same fix as run_graphrag_eval.py: force UTF-8 console output so a
# generated answer with an ordinary non-ASCII character doesn't crash a
# print() and get mistaken for a real pipeline failure.
sys.stdout.reconfigure(encoding="utf-8", errors="replace")
sys.stderr.reconfigure(encoding="utf-8", errors="replace")

INPUT_JSON   = "graphrag_ragas_inputs.json"
OUTPUT_DIR   = Path("ragas_outputs_graphrag")
OUTPUT_XLSX  = OUTPUT_DIR / "aeroops_graphrag_ragas_results.xlsx"
OUTPUT_JSON  = OUTPUT_DIR / "aeroops_graphrag_ragas_results.json"

OUTPUT_DIR.mkdir(parents=True, exist_ok=True)

# ---------------------------------------------------------------------------
# Load secrets so OPENROUTER_API_KEY (and anything else in the file) is
# available via os.environ
# ---------------------------------------------------------------------------
try:
    import tomllib
except ImportError:
    try:
        import tomli as tomllib
    except ImportError:
        tomllib = None

secrets_path = Path(__file__).parent / ".streamlit" / "secrets.toml"
if tomllib and secrets_path.exists():
    with open(secrets_path, "rb") as f:
        for k, v in tomllib.load(f).items():
            os.environ.setdefault(k, str(v))
    print(f"Secrets loaded from {secrets_path}")


def load_inputs(path: str) -> List[Dict[str, Any]]:
    with open(path, encoding="utf-8") as f:
        rows = json.load(f)
    print(f"Loaded {len(rows)} items from {path}")
    return rows


def strip_citations(answer: str) -> str:
    """
    Removes inline citation tags like [D3_c32] or [D1_c15, D3_c27] from an
    answer before it's scored. RAGAS's faithfulness metric decomposes an
    answer into atomic claims and checks each against the retrieved
    evidence text -- a bracketed chunk-ID tag isn't itself a factual claim,
    but if the decomposer treats it as one, no evidence chunk literally
    contains the text "D3_c32" as prose, so that fragment fails grounding
    and drags the whole score down, even when the actual content is
    correct. Verified this against three manually-checked, confirmed-
    correct GraphRAG answers that were all scoring 0.0 faithfulness with
    citations intact. Only strips what goes to the judge -- the saved
    "Model Answer" field keeps citations for actual human/portfolio use.

    Also strips the lenticular-bracket variant (e.g. "【D3_c32】",
    U+3010/U+3011) -- caught 2026-08-30 by directly reading model output
    that used it instead of ASCII brackets. An ASCII-only pattern here
    would leave that tag literally in the text sent to the judge, hitting
    the exact same false-claim mechanism this function exists to prevent.
    """
    return re.sub(r"\s*[\[【][A-Za-z0-9_,\s]+[\]】]", "", answer).strip()


def prepare_contexts(contexts: List[str], max_contexts: int = 16,
                     max_chars: int = 1200) -> List[str]:
    # Fixed 2026-08-30: max_contexts was 15, one below the retriever's own
    # MAX_RECALL_CHUNKS cap of 16 (graphretriever_v5.py) -- so any item that
    # hit the retriever's full budget silently lost its 16th chunk here,
    # before the RAGAS judge ever saw it. Checked live data: 22/45 items
    # were at or over 15 contexts, several far over (OS08 had 25) due to a
    # separate uncapped-merge bug (now fixed in graphretriever_v5.py) --
    # this bound now matches what the retriever actually intends to send.
    # max_chars=1200 was never the issue: chunks here average ~104 chars,
    # max 164, nowhere near that ceiling.
    cleaned = []
    for c in contexts:
        if not c:
            continue
        text = " ".join(str(c).split())
        if len(text) > max_chars:
            text = text[:max_chars].rstrip() + " ..."
        cleaned.append(text)
        if len(cleaned) >= max_contexts:
            break
    return cleaned


def build_ragas_judge():
    openrouter_key = os.getenv("OPENROUTER_API_KEY")
    if not openrouter_key:
        raise RuntimeError("OPENROUTER_API_KEY not set.")

    from langchain_openai import ChatOpenAI
    from ragas.llms import LangchainLLMWrapper

    model = os.getenv("RAGAS_JUDGE_MODEL", "openai/gpt-4o-mini")
    judge = ChatOpenAI(
        model=model,
        api_key=openrouter_key,
        base_url="https://openrouter.ai/api/v1",
        temperature=0,
        timeout=180,
        max_retries=3,
        max_tokens=4000,
        # gpt-4o-mini is NOT a reasoning model -- no internal chain-of-thought
        # burning unpredictable output tokens, no reasoning_effort param to
        # configure or get wrong. Claim verification doesn't need anywhere
        # near its 16,384 ceiling; 4000 is generous headroom for this task.
        # Paid, multi-provider-served on OpenRouter -- not a deprioritized
        # free-tier queue, so no more multi-minute stalls or 504 timeouts.
        # Also the exact model RAGAS's own docs use as their canonical
        # llm_factory example -- well-trodden, not experimental.
        default_headers={
            "HTTP-Referer": "https://palakporwal.site",
            "X-Title": "AeroOps RAGAS Eval",
        },
    )
    print(f"RAGAS judge: OpenRouter / {model}")
    return LangchainLLMWrapper(judge)


def run_ragas(rows: List[Dict[str, Any]]):
    from ragas import evaluate
    from datasets import Dataset

    try:
        from ragas.metrics import faithfulness, context_recall
        metrics = [faithfulness, context_recall]
        dataset = Dataset.from_list([
            {
                "question":     r["Question"],
                "answer":       strip_citations(r["Model Answer"]),
                "contexts":     prepare_contexts(r.get("Retrieved Contexts") or []),
                "ground_truth": r["Ground Truth"],
            }
            for r in rows
        ])
        schema = "legacy"
    except Exception:
        from ragas.metrics import Faithfulness, LLMContextRecall
        metrics = [Faithfulness(), LLMContextRecall()]
        dataset = Dataset.from_list([
            {
                "user_input":          r["Question"],
                "response":            strip_citations(r["Model Answer"]),
                "retrieved_contexts":  prepare_contexts(r.get("Retrieved Contexts") or []),
                "reference":           r["Ground Truth"],
            }
            for r in rows
        ])
        schema = "new"

    print(f"Running RAGAS ({schema} schema) on {len(rows)} items...")

    try:
        from ragas.run_config import RunConfig
        run_config = RunConfig(max_workers=1, timeout=180, max_retries=3)
    except Exception:
        run_config = None

    judge = build_ragas_judge()

    kwargs = {"dataset": dataset, "metrics": metrics,
              "llm": judge, "raise_exceptions": False}
    if run_config:
        kwargs["run_config"] = run_config

    return evaluate(**kwargs)


def merge_scores(ragas_result, rows):
    try:
        df = ragas_result.to_pandas()
        scores = df.to_dict(orient="records")
    except Exception:
        scores = [{} for _ in rows]

    merged = []
    for i, row in enumerate(rows):
        s = scores[i] if i < len(scores) else {}
        # BUG FIX: `s.get("faithfulness") or s.get("Faithfulness")` silently
        # discarded genuine 0.0 scores, since 0.0 is falsy in Python and
        # `or` falls through to the next operand whenever the first is
        # falsy. A real, computed 0.0 got treated as "nothing here" and
        # replaced with s.get("Faithfulness") (capital F, a key that never
        # exists in RAGAS's actual output), silently becoming None. This
        # is very likely the true explanation for every "missing" score
        # chased across tonight's retries -- they weren't API failures,
        # they were real 0.0 judgments this line was throwing away.
        f_val = s.get("faithfulness")
        if f_val is None:
            f_val = s.get("Faithfulness")
        r_val = s.get("context_recall")
        if r_val is None:
            r_val = s.get("LLMContextRecall")
        merged.append({
            **row,
            "Faithfulness":    f_val,
            "Context Recall":  r_val,
        })
    return merged


def save_excel(rows, path):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment

    wb = Workbook()
    ws = wb.active
    ws.title = "GraphRAG RAGAS"

    headers = ["ID", "Category", "Question", "Expected Answer",
               "Model Answer", "Faithfulness", "Context Recall", "Time (s)"]

    hfont = Font(bold=True, color="FFFFFF")
    hfill = PatternFill("solid", fgColor="2F5496")
    for col, h in enumerate(headers, 1):
        c = ws.cell(row=1, column=col, value=h)
        c.font = hfont
        c.fill = hfill
        c.alignment = Alignment(horizontal="center", wrap_text=True)

    for i, row in enumerate(rows):
        ws.append([
            row.get("ID"), row.get("Category"), row.get("Question"),
            row.get("Ground Truth"), row.get("Model Answer"),
            row.get("Faithfulness"), row.get("Context Recall"),
            row.get("Time (s)"),
        ])
        for cell in ws[i + 2]:
            cell.alignment = Alignment(wrap_text=True, vertical="top")

    for col, width in zip("ABCDEFGH", [8, 18, 50, 50, 50, 14, 14, 10]):
        ws.column_dimensions[col].width = width

    ws.freeze_panes = "A2"
    wb.save(path)
    print(f"Saved Excel -> {path}")


def print_summary(rows):
    # BUG FIX: `r.get(key) is not None` doesn't catch RAGAS's own
    # float('nan') return value for a metric job that failed at the API
    # level (raise_exceptions=False means those come back as NaN, not
    # None, not an exception). NaN passed `is not None` and went into
    # `vals`, and statistics.mean() of any list containing NaN returns
    # NaN -- so ONE item's transient API failure (confirmed: CD08 hit an
    # APIConnectionError/TimeoutError on both its Faithfulness and
    # Context Recall judge calls) silently zeroed out the mean for all
    # 45 items, discarding 44 perfectly good scores. This is a different
    # bug from the truthiness bug documented in merge_scores() above --
    # that one was about 0.0 becoming None; this one is about NaN
    # skipping the None filter entirely.
    def mean(key):
        vals = [float(r[key]) for r in rows
                if r.get(key) is not None and not math.isnan(float(r[key]))]
        skipped = sum(1 for r in rows if r.get(key) is None
                      or (isinstance(r.get(key), float) and math.isnan(r[key])))
        if skipped:
            print(f"  ({skipped} item(s) excluded from {key} mean -- "
                  f"missing/NaN score, not a real 0)")
        return round(statistics.mean(vals), 3) if vals else "N/A"

    print("\n======= GRAPHRAG RAGAS SUMMARY =======")
    print(f"Items evaluated: {len(rows)}")
    print(f"Faithfulness mean:   {mean('Faithfulness')}")
    print(f"Context recall mean: {mean('Context Recall')}")
    print("========================================\n")


def main():
    rows = load_inputs(INPUT_JSON)
    ragas_result = run_ragas(rows)
    rows_scored = merge_scores(ragas_result, rows)

    save_excel(rows_scored, OUTPUT_XLSX)

    with open(OUTPUT_JSON, "w", encoding="utf-8") as f:
        json.dump(rows_scored, f, indent=2, ensure_ascii=False,
                  default=str)
    print(f"Saved JSON  -> {OUTPUT_JSON}")

    print_summary(rows_scored)


if __name__ == "__main__":
    main()
