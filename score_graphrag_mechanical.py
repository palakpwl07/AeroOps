# score_graphrag_mechanical.py
#
# Deterministic, graph-native scoring for GraphRAG -- no LLM judge, no
# prose interpretation, nothing to game. Reads graphrag_ragas_inputs.json
# (already produced by run_graphrag_eval.py) and the frozen gold set, and
# checks the SYSTEM's actual retrieval/reasoning directly against known-
# correct graph node ids, chunk ids, and expected causal chains.
#
# This was already planned for -- run_graphrag_eval.py's own comment says
# it kept Matched Entities/Path Nodes/Path Edges/Reasoning Chain
# "so the later mechanical scoring layer (node/edge recall against
# ground_truth_nodes, path-order against expected_chain, citation
# precision) doesn't require rerunning the pipeline." That layer was
# never built until now.
#
# Metrics (all GraphRAG-only -- naive RAG has no graph structure to check
# node/edge/path metrics against; only citation precision/recall has a
# naive-RAG analogue, and this script doesn't compute that side):
#   1. Node recall/precision   -- matched_entities + path nodes vs
#                                  ground_truth_nodes
#   2. Citation recall/precision -- chunk IDs cited in the answer vs
#                                  gold source_chunk_ids
#   3. Chain coverage           -- for items with a non-empty
#                                  expected_chain, what fraction of the
#                                  expected chunk sequence appears
#                                  anywhere in the system's actual
#                                  reasoning chain (order-insensitive;
#                                  a stricter exact-order match is also
#                                  reported)
#   4. Edge authenticity rate   -- of the reasoning-chain hops that carry
#                                  an edge_type, what fraction are real
#                                  claim-backed edges (CAUSES, LEADS_TO,
#                                  etc, all with a claim_id) vs
#                                  structural/co-occurrence hops
#                                  (MENTIONED_IN, HAS_CHUNK, or any edge
#                                  missing a claim_id)
#   5. Over-fetch ratio         -- contexts retrieved vs gold chunks
#                                  actually needed
#
# Usage:
#   python score_graphrag_mechanical.py
#
# Output (deliberately its own folder, not ragas_outputs_graphrag/ --
# that one's already crowded with two RAGAS runs plus patched variants):
#   graphrag_mechanical_scores/aeroops_graphrag_mechanical_scores.json
#   graphrag_mechanical_scores/aeroops_graphrag_mechanical_scores.xlsx

from __future__ import annotations

import json
import re
import statistics
import sys
from pathlib import Path
from typing import Any, Dict, List, Optional, Set

sys.stdout.reconfigure(encoding="utf-8", errors="replace")

INPUT_JSON = "graphrag_ragas_inputs.json"
GOLD_JSON = "aeroops_goldset_v1_frozen.json"
OUTPUT_DIR = Path("graphrag_mechanical_scores")
OUTPUT_JSON = OUTPUT_DIR / "aeroops_graphrag_mechanical_scores.json"
OUTPUT_XLSX = OUTPUT_DIR / "aeroops_graphrag_mechanical_scores.xlsx"

OUTPUT_DIR.mkdir(parents=True, exist_ok=True)

# Same set as context_builder.py's REAL_EDGE_TYPES -- edges the retriever's
# star queries populate with a claim_id (extracted/verified from a
# document), as opposed to MENTIONED_IN/HAS_CHUNK (structural,
# no claim_id, Entity->Chunk co-occurrence rather than an asserted
# relationship between two entities).
REAL_EDGE_TYPES = {
    "MANIFESTS_AS", "CAUSES", "INFLUENCES", "LEADS_TO", "MITIGATES",
    "AFFECTS", "DEGRADES", "RESTORES", "INDICATES", "DETECTS",
}

# Fixed 2026-08-30: the model sometimes cites with lenticular brackets
# (e.g. "【D3_c03】", U+3010/U+3011) instead of ASCII "[...]" -- confirmed
# by directly reading its actual output, not assumed. An ASCII-only
# regex silently scored those citations as absent, producing false
# "citation_recall got worse" diffs that were really just bracket-style
# variance with the same chunk ID correctly cited underneath.
CITATION_RE = re.compile(r"[\[【]([A-Za-z0-9_,\s]+)[\]】]")


def load_inputs() -> List[Dict[str, Any]]:
    with open(INPUT_JSON, encoding="utf-8") as f:
        return json.load(f)


def load_gold() -> Dict[str, Dict[str, Any]]:
    with open(GOLD_JSON, encoding="utf-8") as f:
        gold = json.load(f)
    return {g["id"]: g for g in gold}


def extract_cited_chunks(answer: str) -> Set[str]:
    """Every bracketed chunk-id tag the answer actually cites, e.g.
    '...tailpipe fires. [D3_c22]' or '[D1_c15, D3_c27]'."""
    cited: Set[str] = set()
    for bracket in CITATION_RE.findall(answer or ""):
        for cid in bracket.split(","):
            cid = cid.strip()
            if re.match(r"^D\d+_c\d+$", cid):
                cited.add(cid)
    return cited


def safe_ratio(numerator: int, denominator: int) -> Optional[float]:
    if denominator == 0:
        return None
    return round(numerator / denominator, 3)


def score_node_coverage(row: Dict[str, Any], gt_nodes: Set[str]) -> Dict[str, Any]:
    matched_ids = {m.get("id") for m in (row.get("Matched Entities") or []) if m.get("id")}
    path_node_ids = {n.get("id") for n in (row.get("Path Nodes") or []) if n.get("id")}
    chain_ids = {s.get("id") for s in (row.get("Reasoning Chain") or []) if s.get("id")}
    touched = matched_ids | path_node_ids | chain_ids

    if not gt_nodes:
        return {"node_recall": None, "node_precision": None,
                "touched_nodes": sorted(touched), "missing_nodes": []}

    overlap = touched & gt_nodes
    return {
        "node_recall": safe_ratio(len(overlap), len(gt_nodes)),
        "node_precision": safe_ratio(len(overlap), len(touched)),
        "touched_nodes": sorted(touched),
        "missing_nodes": sorted(gt_nodes - touched),
    }


def score_citations(row: Dict[str, Any], gt_chunks: Set[str]) -> Dict[str, Any]:
    cited = extract_cited_chunks(row.get("Model Answer", ""))
    if not gt_chunks:
        return {"citation_recall": None, "citation_precision": None,
                "cited_chunks": sorted(cited), "missing_chunks": []}
    overlap = cited & gt_chunks
    return {
        "citation_recall": safe_ratio(len(overlap), len(gt_chunks)),
        "citation_precision": safe_ratio(len(overlap), len(cited)),
        "cited_chunks": sorted(cited),
        "missing_chunks": sorted(gt_chunks - cited),
    }


def score_chain_coverage(row: Dict[str, Any], expected_chain: List[Dict[str, Any]]) -> Dict[str, Any]:
    if not expected_chain:
        return {"chain_coverage": None, "chain_exact_order": None}

    expected_seq = [step["source_chunk_id"] for step in expected_chain if step.get("source_chunk_id")]
    if not expected_seq:
        return {"chain_coverage": None, "chain_exact_order": None}

    actual_seq: List[str] = []
    for step in (row.get("Reasoning Chain") or []):
        for cid in (step.get("edge_chunks") or []):
            if cid not in actual_seq:
                actual_seq.append(cid)

    covered = [c for c in expected_seq if c in actual_seq]
    return {
        "chain_coverage": safe_ratio(len(covered), len(expected_seq)),
        "chain_exact_order": actual_seq[:len(expected_seq)] == expected_seq,
        "expected_sequence": expected_seq,
        "actual_sequence": actual_seq,
    }


def score_edge_authenticity(row: Dict[str, Any]) -> Dict[str, Any]:
    chain = row.get("Reasoning Chain") or []
    hops = [s for s in chain if s.get("edge_type")]
    if not hops:
        return {"edge_authenticity_rate": None, "real_hops": 0,
                "structural_hops": 0, "structural_edge_types": []}

    real = 0
    structural_types = []
    for hop in hops:
        is_real = hop["edge_type"] in REAL_EDGE_TYPES and bool(hop.get("edge_claim_id"))
        if is_real:
            real += 1
        else:
            structural_types.append(hop["edge_type"])

    return {
        "edge_authenticity_rate": safe_ratio(real, len(hops)),
        "real_hops": real,
        "structural_hops": len(hops) - real,
        "structural_edge_types": structural_types,
    }


def score_overfetch(row: Dict[str, Any], gt_chunks: Set[str]) -> Optional[float]:
    n_contexts = len(row.get("Retrieved Contexts") or [])
    if not gt_chunks:
        return None
    return round(n_contexts / len(gt_chunks), 2)


def main():
    rows = load_inputs()
    gold = load_gold()
    print(f"Loaded {len(rows)} items from {INPUT_JSON}")

    scored = []
    for row in rows:
        qid = row["ID"]
        g = gold.get(qid, {})
        gt_nodes = set(g.get("ground_truth_nodes") or [])
        gt_chunks = set(c.strip() for c in (g.get("source_chunk_ids") or []))
        expected_chain = g.get("expected_chain") or []

        item = {
            "ID": qid,
            "Category": row.get("Category"),
            **score_node_coverage(row, gt_nodes),
            **score_citations(row, gt_chunks),
            **score_chain_coverage(row, expected_chain),
            **score_edge_authenticity(row),
            "overfetch_ratio": score_overfetch(row, gt_chunks),
            "num_contexts": len(row.get("Retrieved Contexts") or []),
        }
        scored.append(item)

    # ---- Category + overall summary ----
    def mean_of(key, items):
        vals = [it[key] for it in items if it.get(key) is not None]
        return round(statistics.mean(vals), 3) if vals else None

    categories = sorted({it["Category"] for it in scored if it.get("Category")})
    metric_keys = ["node_recall", "node_precision", "citation_recall",
                   "citation_precision", "chain_coverage",
                   "edge_authenticity_rate", "overfetch_ratio"]

    summary = {"overall": {k: mean_of(k, scored) for k in metric_keys}}
    summary["overall"]["n_items"] = len(scored)
    for cat in categories:
        cat_items = [it for it in scored if it["Category"] == cat]
        summary[cat] = {k: mean_of(k, cat_items) for k in metric_keys}
        summary[cat]["n_items"] = len(cat_items)

    print("\n======= GRAPHRAG MECHANICAL SCORING SUMMARY =======")
    header = f"{'Category':22s}" + "".join(f"{k[:14]:>16s}" for k in metric_keys)
    print(header)
    for cat in ["overall"] + categories:
        row_vals = summary[cat]
        line = f"{cat:22s}" + "".join(
            f"{('N/A' if row_vals.get(k) is None else row_vals[k]):>16}" for k in metric_keys
        )
        print(line)
    print("====================================================\n")

    output = {"summary": summary, "items": scored}
    with open(OUTPUT_JSON, "w", encoding="utf-8") as f:
        json.dump(output, f, indent=2, ensure_ascii=False)
    print(f"Saved JSON -> {OUTPUT_JSON}")

    # ---- Excel ----
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment

    wb = Workbook()

    ws1 = wb.active
    ws1.title = "Summary"
    ws1.append(["Category", "N Items"] + metric_keys)
    for col, h in enumerate(["Category", "N Items"] + metric_keys, 1):
        c = ws1.cell(row=1, column=col, value=h)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="2F5496")
    for cat in ["overall"] + categories:
        row_vals = summary[cat]
        ws1.append([cat, row_vals["n_items"]] + [row_vals.get(k) for k in metric_keys])
    for col_letter in "ABCDEFGHI":
        ws1.column_dimensions[col_letter].width = 18

    ws2 = wb.create_sheet("Per-Item")
    item_headers = ["ID", "Category", "node_recall", "node_precision",
                     "citation_recall", "citation_precision",
                     "chain_coverage", "chain_exact_order",
                     "edge_authenticity_rate", "real_hops", "structural_hops",
                     "overfetch_ratio", "num_contexts", "missing_nodes",
                     "missing_chunks", "structural_edge_types"]
    ws2.append(item_headers)
    for col, h in enumerate(item_headers, 1):
        c = ws2.cell(row=1, column=col, value=h)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="2F5496")
    for it in scored:
        ws2.append([
            it["ID"], it["Category"], it.get("node_recall"), it.get("node_precision"),
            it.get("citation_recall"), it.get("citation_precision"),
            it.get("chain_coverage"), it.get("chain_exact_order"),
            it.get("edge_authenticity_rate"), it.get("real_hops"), it.get("structural_hops"),
            it.get("overfetch_ratio"), it.get("num_contexts"),
            ", ".join(it.get("missing_nodes") or []),
            ", ".join(it.get("missing_chunks") or []),
            ", ".join(it.get("structural_edge_types") or []),
        ])
    for col_letter, width in zip("ABCDEFGHIJKLMNOP", [8, 18, 12, 12, 12, 12, 12, 12, 14, 10, 12, 12, 12, 24, 24, 24]):
        ws2.column_dimensions[col_letter].width = width
    ws2.freeze_panes = "A2"

    wb.save(OUTPUT_XLSX)
    print(f"Saved Excel -> {OUTPUT_XLSX}")


if __name__ == "__main__":
    main()
